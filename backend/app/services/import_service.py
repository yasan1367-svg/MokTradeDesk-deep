from datetime import datetime
from sqlalchemy.orm import Session
from typing import List, Dict, Any, Optional
from bs4 import BeautifulSoup

from ..models.strategy import Trade, TradeSource, TestType
from ..utils.trade_metrics import calculate_r_multiple


# ═════════════════════════════════════════════
# Soft4X Importer
# ═════════════════════════════════════════════
class Soft4XImporter:
    """واردکننده فایل‌های اکسل خروجی Soft4X"""

    def __init__(self, db: Session, symbol: str = "XAUUSD", test_type: str = "backtest"):
        self.db = db
        self.symbol = symbol
        self.test_type = TestType(test_type)

    def parse_file(self, file_path: str) -> List[Dict[str, Any]]:
        from openpyxl import load_workbook

        wb = load_workbook(file_path, data_only=True)
        ws = wb["Trades"] if "Trades" in wb.sheetnames else wb.active

        headers = []
        for cell in ws[1]:
            headers.append(cell.value)

        col_index = {}
        for idx, header in enumerate(headers):
            if header:
                col_index[str(header).strip()] = idx

        trades = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue

            if col_index.get("Open Time") is None:
                continue

            open_time = self._to_datetime(self._get_value(row, col_index.get("Open Time")))
            close_time = self._to_datetime(self._get_value(row, col_index.get("Close Time")))

            if open_time is None:
                continue

            direction = self._get_direction(self._get_value(row, col_index.get("Type")))
            open_price = float(self._get_value(row, col_index.get("Open Price"), 0) or 0)
            close_price = float(self._get_value(row, col_index.get("Close Price"), 0) or 0)
            sl = self._to_float(self._get_value(row, col_index.get("SL")))

            trade = {
                "symbol": self.symbol,
                "test_type": self.test_type,
                "direction": direction,
                "open_time": open_time,
                "close_time": close_time,
                "open_price": open_price,
                "close_price": close_price,
                "size": float(self._get_value(row, col_index.get("Size"), 0) or 0),
                "sl": sl,
                "tp": self._to_float(self._get_value(row, col_index.get("TP"))),
                "pnl": float(self._get_value(row, col_index.get("P/L"), 0) or 0),
                "r_multiple": calculate_r_multiple(direction, open_price, close_price, sl),
                "commission": self._to_float(self._get_value(row, col_index.get("Commission"))) or 0,
                "swap": 0.0,
                "entry_sequence": 1,
                "source": TradeSource.SOFT4X_IMPORT,
                "raw_data": self._make_json_safe({
                    headers[i]: row[i] for i in range(len(row)) if i < len(headers)
                })
            }
            trades.append(trade)

        return trades

    def save_trades(
        self,
        trades: List[Dict[str, Any]],
        version_id: Optional[int] = None,
        prop_stage_id: Optional[int] = None,
    ) -> List[Trade]:
        db_trades = []
        for trade_data in trades:
            db_trade = Trade(
                version_id=version_id,
                prop_stage_id=prop_stage_id,
                **trade_data
            )
            self.db.add(db_trade)
            db_trades.append(db_trade)

        self.db.commit()

        if prop_stage_id:
            self._update_prop_stage_profit(prop_stage_id)

        return db_trades

    def _update_prop_stage_profit(self, prop_stage_id: int):
        from ..models.prop import PropStage, StageType

        stage = self.db.query(PropStage).filter(PropStage.id == prop_stage_id).first()
        if not stage:
            return

        all_trades = self.db.query(Trade).filter(Trade.prop_stage_id == prop_stage_id).all()
        total_pnl = sum(t.pnl or 0 for t in all_trades)

        if stage.stage_type == StageType.FUNDED_REAL:
            share = (stage.profit_share_percentage or 80.0) / 100.0
            stage.current_profit = total_pnl * share
        else:
            stage.current_profit = total_pnl

        self.db.commit()

    def _apply_symbol_mapping(self, symbol: str) -> str:
        """تبدیل نماد اصلی به نماد استاندارد"""
        from ..models.strategy import SymbolMapping
        mapping = self.db.query(SymbolMapping).filter(
            SymbolMapping.original_symbol == symbol
        ).first()
        if mapping:
            return mapping.canonical_symbol
        return symbol

    def _get_value(self, row, index, default=None):
        if index is None:
            return default
        try:
            value = row[index] if index < len(row) else default
            return value if value is not None else default
        except:
            return default

    def _get_direction(self, value):
        if value:
            return "buy" if str(value).lower() == "buy" else "sell"
        return "sell"

    def _to_float(self, value):
        if value is None:
            return None
        try:
            return float(value)
        except:
            return None

    def _to_datetime(self, value):
        if value is None:
            return None
        if isinstance(value, datetime):
            return value
        try:
            if hasattr(value, 'to_pydatetime'):
                return value.to_pydatetime()
            return datetime.strptime(str(value), "%Y-%m-%d %H:%M:%S")
        except:
            try:
                return datetime.fromisoformat(str(value))
            except:
                return None

    def _make_json_safe(self, data: dict) -> dict:
        safe_data = {}
        for key, value in data.items():
            if value is None:
                safe_data[key] = None
            elif isinstance(value, datetime):
                safe_data[key] = value.isoformat()
            elif hasattr(value, 'to_pydatetime'):
                safe_data[key] = value.to_pydatetime().isoformat()
            elif hasattr(value, 'isoformat'):
                safe_data[key] = value.isoformat()
            elif isinstance(value, (int, float, str, bool)):
                safe_data[key] = value
            else:
                safe_data[key] = str(value)
        return safe_data


# ═════════════════════════════════════════════
# MT4 Importer
# ═════════════════════════════════════════════
class MT4Importer:
    """واردکننده فایل‌های HTML متاتریدر (فقط بخش Positions)"""

    def __init__(self, db: Session, test_type: str = "backtest"):
        self.db = db
        self.test_type = TestType(test_type)

    def parse_html(self, html_content: str) -> List[Dict[str, Any]]:
        soup = BeautifulSoup(html_content, 'html.parser')
        trades = []
        rows = soup.find_all('tr')

        in_positions_section = False
        header_skipped = False

        for idx, row in enumerate(rows):
            row_text = row.get_text(strip=True)

            if 'Positions' in row_text and row.find('th'):
                in_positions_section = True
                header_skipped = False
                continue

            if ('Orders' in row_text or 'Deals' in row_text) and row.find('th'):
                in_positions_section = False
                continue

            if not in_positions_section:
                continue

            if not header_skipped:
                header_skipped = True
                continue

            all_cells = row.find_all('td')
            visible_cells = [
                c for c in all_cells
                if 'hidden' not in (c.get('class') or [])
            ]

            if len(visible_cells) >= 13:
                try:
                    trade = self._parse_row(visible_cells)
                    if trade:
                        trades.append(trade)
                except Exception as e:
                    print(f"  ❌ خطا: {e}")

        return trades

    def _parse_row(self, cells) -> Optional[Dict[str, Any]]:
        try:
            open_time = self._to_datetime(cells[0].get_text(strip=True))
            position = cells[1].get_text(strip=True)
            symbol = cells[2].get_text(strip=True)
            direction_raw = cells[3].get_text(strip=True).lower()
            volume = float(cells[4].get_text(strip=True))
            open_price = float(cells[5].get_text(strip=True).replace(',', ''))
            sl = self._to_float(cells[6].get_text(strip=True))
            tp = self._to_float(cells[7].get_text(strip=True))
            close_time = self._to_datetime(cells[8].get_text(strip=True))
            close_price = self._to_float(cells[9].get_text(strip=True))
            commission = self._to_float(cells[10].get_text(strip=True)) or 0
            swap = self._to_float(cells[11].get_text(strip=True)) or 0
            profit = self._to_float(cells[12].get_text(strip=True)) or 0

            if open_time is None:
                return None

            direction = "buy" if "buy" in direction_raw else "sell"

            return {
                "symbol": self._apply_symbol_mapping(symbol),
                "test_type": self.test_type,
                "direction": direction,
                "open_time": open_time,
                "close_time": close_time,
                "open_price": open_price,
                "close_price": close_price,
                "size": volume,
                "sl": sl,
                "tp": tp,
                "pnl": profit,
                "r_multiple": calculate_r_multiple(direction, open_price, close_price, sl),
                "commission": commission,
                "swap": swap,
                "entry_sequence": 1,
                "source": TradeSource.MT4_IMPORT,
                "raw_data": {
                    "position": position,
                    "raw_direction": direction_raw,
                },
            }
        except Exception as e:
            print(f"  ❌ خطای _parse_row: {e}")
            return None

    def save_trades(
        self,
        trades: List[Dict[str, Any]],
        version_id: Optional[int] = None,
        prop_stage_id: Optional[int] = None,
    ) -> List[Trade]:
        db_trades = []
        for trade_data in trades:
            db_trade = Trade(
                version_id=version_id,
                prop_stage_id=prop_stage_id,
                **trade_data
            )
            self.db.add(db_trade)
            db_trades.append(db_trade)

        self.db.commit()

        if prop_stage_id:
            self._update_prop_stage_profit(prop_stage_id)

        return db_trades

    def _update_prop_stage_profit(self, prop_stage_id: int):
        from ..models.prop import PropStage, StageType

        stage = self.db.query(PropStage).filter(PropStage.id == prop_stage_id).first()
        if not stage:
            return

        all_trades = self.db.query(Trade).filter(Trade.prop_stage_id == prop_stage_id).all()
        total_pnl = sum(t.pnl or 0 for t in all_trades)

        if stage.stage_type == StageType.FUNDED_REAL:
            share = (stage.profit_share_percentage or 80.0) / 100.0
            stage.current_profit = total_pnl * share
        else:
            stage.current_profit = total_pnl

        self.db.commit()

    def _apply_symbol_mapping(self, symbol: str) -> str:
        """تبدیل نماد اصلی به نماد استاندارد"""
        from ..models.strategy import SymbolMapping
        mapping = self.db.query(SymbolMapping).filter(
            SymbolMapping.original_symbol == symbol
        ).first()
        if mapping:
            return mapping.canonical_symbol
        return symbol

    def _to_float(self, value: str) -> Optional[float]:
        if not value or value in ['', '-', 'N/A']:
            return None
        try:
            return float(value.replace(',', '').replace(' ', ''))
        except:
            return None

    def _to_datetime(self, value: str) -> Optional[datetime]:
        if not value:
            return None
        try:
            return datetime.strptime(value.strip(), "%Y.%m.%d %H:%M:%S")
        except:
            try:
                return datetime.strptime(value.strip(), "%Y.%m.%d %H:%M")
            except:
                return None